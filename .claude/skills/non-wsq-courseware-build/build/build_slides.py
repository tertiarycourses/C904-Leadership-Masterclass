#!/usr/bin/env python3
"""C904 visual deck from the shared course metadata and original concept content."""
from pathlib import Path
import sys,os,re,math
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
import course_data as C
root=Path(os.environ['COURSE_REPO']); prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
INK='152A3A'; TEAL='087F8C'; GOLD='B57D28'; GREY='526474'; LIGHT='EFF5F6'; WHITE='FFFFFF'; NAVY='143744'
accents=['087F8C','245A81','885E34','42734A','835165','506480']
def color(c):return RGBColor.from_string(c)
def shape(s,x,y,w,h,c,kind=MSO_SHAPE.RECTANGLE):
 z=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));z.fill.solid();z.fill.fore_color.rgb=color(c);z.line.fill.background();return z
def text(s,x,y,w,h,txt,size=22,c=INK,bold=False,align=PP_ALIGN.LEFT):
 z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=z.text_frame;f.word_wrap=True;f.margin_left=0;f.margin_right=0;f.margin_top=0
 for i,line in enumerate(txt.split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.alignment=align;p.space_after=Pt(8);r=p.add_run();r.text=line;r.font.name='Aptos';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color(c)
 return z
def base(title,kicker='C904 · LEADERSHIP MASTERCLASS',accent=TEAL):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(WHITE)
 text(s,.6,.35,12,.3,kicker.upper(),11,accent,True)
 title=title.strip(); sz=32 if len(title)<65 else 27
 text(s,.6,.9,12.1,1.1,title,sz,INK,True)
 text(s,.6,7.12,9,.2,'Leadership Masterclass · C904 · v4.0 | Tertiary Infotech Academy',9,GREY)
 text(s,12.1,7.1,.6,.25,str(len(prs.slides)),10,GREY,align=PP_ALIGN.RIGHT)
 return s
def cards(title,items,kicker='',accent=TEAL):
    # Keep a short source label with its explanation and balance continuation pages.
    merged=[];i=0
    while i<len(items):
        t=items[i]
        if len(t)<55 and i+1<len(items) and len(items[i+1])>100 and (t.endswith(':') or not t.endswith('.')):
            t=t.rstrip(':')+': '+items[i+1];i+=1
        merged.append(t);i+=1
    items=[]
    for para in merged:
        if len(para)<=240: items.append(para); continue
        sentences=re.split(r'(?<=[.!?])\s+',para);buf=''
        for sentence in sentences:
            if buf and len(buf)+len(sentence)>240:items.append(buf);buf=''
            if len(sentence)>260:
                words=sentence.split();frag=''
                for word in words:
                    if len(frag)+len(word)>230:items.append(frag);frag=''
                    frag+=(" " if frag else "")+word
                if frag:items.append(frag)
            else:buf+=(" " if buf else "")+sentence
        if buf:items.append(buf)
    maxn=6 if all(len(x)<=45 for x in items) and not any("http" in x for x in items) else 4
    pages=max(1,math.ceil(len(items)/maxn));chunk=math.ceil(len(items)/pages)
    for q in range(pages):
        part=items[q*chunk:(q+1)*chunk];s=base(title+(' — continued' if q else ''),kicker or 'CORE CONCEPT',accent);n=len(part)
        if n>=4 and all(len(x)<=45 for x in part) and 'http' not in ' '.join(part):
            # A hub-and-spoke framework for related traits, effects or choices.
            cx,cy=6.66,4.5
            shape(s,cx-1.05,cy-.65,2.1,1.3,accent,MSO_SHAPE.OVAL)
            text(s,cx-.9,cy-.35,1.8,.8,'Leadership\nin action',20,WHITE,True,PP_ALIGN.CENTER)
            for i,t in enumerate(part):
                angle=2*math.pi*i/n-math.pi/2;x=cx+4.22*math.cos(angle)-1.75;y=cy+1.75*math.sin(angle)-.55
                line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(cx),Inches(cy),Inches(x+1.75),Inches(y+.55));line.line.color.rgb=color(accent);line.line.width=Pt(1.5)
                shape(s,x,y,3.5,1.1,LIGHT,MSO_SHAPE.ROUNDED_RECTANGLE);text(s,x+.15,y+.12,3.2,.87,t,18,INK,True,PP_ALIGN.CENTER)
        elif n==1 and len(part[0])<180:
            text(s,.75,2.55,8.2,2.6,part[0],32,accent,True)
            shape(s,9.6,2.45,2.45,2.45,LIGHT,MSO_SHAPE.OVAL)
            shape(s,10.25,3.1,1.1,1.1,accent,MSO_SHAPE.CHEVRON)
            text(s,.75,5.85,10.7,.6,'Discuss the behaviour and its effect in a real work situation.',18,GREY)
        else:
            cols=3 if n==3 and all(len(x)<360 for x in part) else (2 if n>1 else 1)
            rows=math.ceil(n/cols);w=(12.1-.35*(cols-1))/cols;h=(4.65-.3*(rows-1))/rows
            for i,t in enumerate(part):
                x=.6+(i%cols)*(w+.35);y=2.1+(i//cols)*(h+.3);shape(s,x,y,w,h,LIGHT)
                text(s,x+.23,y+.15,.55,.6,str(q*chunk+i+1).zfill(2),25,accent,True)
                sz=19 if len(t)<170 else 16
                if cols==3: text(s,x+.23,y+.95,w-.46,h-1.1,t,sz)
                else: text(s,x+.95,y+.23,w-1.18,h-.46,t,sz)
    return s
def flow(title,nodes,caption,accent=TEAL,kicker='FRAMEWORK'):
 s=base(title,kicker,accent);n=len(nodes);w=(12.1-.38*(n-1))/n
 for i,label in enumerate(nodes):
  x=.6+i*(w+.38);shape(s,x,2.55,w,1.6,accent,MSO_SHAPE.ROUNDED_RECTANGLE)
  text(s,x+.18,2.75,w-.36,1.15,label,22,WHITE,True,PP_ALIGN.CENTER)
  if i<n-1:shape(s,x+w+.06,3.1,.24,.35,GOLD,MSO_SHAPE.CHEVRON)
 text(s,.85,4.75,11.6,1.65,caption,22)
 return s
def divider(n,title,subtitle):
 s=base(title,f'TOPIC {n:02d}',accents[n-1]);shape(s,8.7,2.2,3.9,3.9,LIGHT,MSO_SHAPE.OVAL);text(s,9.05,2.75,3.2,2,str(n).zfill(2),100,accents[n-1],True,PP_ALIGN.CENTER)
 text(s,.65,2.55,7.4,1.7,subtitle,29);text(s,.65,5.2,7.4,.7,'Understand → discuss → practise → reflect',20,GREY)
def theory_visual(o,accent):
    low=o['title'].lower();s=base(o['title'],'MOTIVATION MODEL · DISCUSSION LENS',accent)
    if 'maslow' in low:
        labels=['Self-actualisation','Esteem','Belonging','Safety','Physiological needs']
        for j,label in enumerate(labels):
            w=3.0+j*.85;x=3.7-w/2;y=2.2+j*.72
            shape(s,x,y,w,.63,accent,MSO_SHAPE.TRAPEZOID);text(s,x+.1,y+.15,w-.2,.35,label,16,WHITE,True,PP_ALIGN.CENTER)
        text(s,7.7,2.5,4.8,3.6,'Needs can overlap and vary by context.\n\nAsk which work conditions or needs require attention. Do not treat the hierarchy as a fixed sequence for every person.',24)
    elif 'herzberg' in low:
        for x,label,body in [(0.7,'Hygiene factors','Conditions, policies and supervision\n\nAddress sources of dissatisfaction.'),(7.0,'Motivators','Achievement, responsibility and meaningful work\n\nExplore opportunities for satisfaction and growth.')]:
            shape(s,x,2.3,5.6,3.5,LIGHT);text(s,x+.3,2.65,5,1,label,30,accent,True);text(s,x+.3,3.65,5,2.05,body,20)
        text(s,.7,6.2,11.8,.55,'Check the person and context; use the distinction as a lens rather than a prediction.',19,GREY)
    else:
        for x,label,body in [(0.7,'Autonomy','Choice within clear boundaries'),(4.9,'Competence','Effective action, practice and growth'),(9.1,'Relatedness','Connection and belonging')]:
            shape(s,x,2.35,3.5,1.7,accent,MSO_SHAPE.OVAL);text(s,x+.2,2.97,3.1,.7,label,24,WHITE,True,PP_ALIGN.CENTER);text(s,x,4.55,3.5,1.65,body,24,INK,True,PP_ALIGN.CENTER)
        text(s,.7,6.5,11.8,.35,'Ask what would help this person experience each need in the work.',19,GREY)
def chart_matrix():
 s=base('A decision matrix makes assumptions visible','FICTIONAL TEACHING DATA');d=CategoryChartData();d.categories=['Quality','Cost','Resilience'];d.add_series('Option A',[4,3,4]);d.add_series('Option B',[3,5,2]);ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.65), Inches(2.15), Inches(7), Inches(4.55),d).chart;ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.value_axis.maximum_scale=5;ch.value_axis.minimum_scale=0
 text(s,8.1,2.3,4.55,1.5,'Weights: 0.4 / 0.3 / 0.3\nA = 3.7     B = 3.3',24,TEAL,True)
 text(s,8.1,4.05,4.55,2.3,'Change weights to 0.2 / 0.6 / 0.2:\nA = 3.4     B = 4.0\nThe ranking changes. Explain why.',22)
# Cover
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(WHITE)
logo=Path(__file__).parent.parent/'assets/tertiary-infotech-logo.png';s.shapes.add_picture(str(logo),Inches(.65),Inches(.5),height=Inches(.65))
text(s,.65,1.85,8.7,1.65,'Leadership\nMasterclass',46,INK,True)
text(s,.65,4.15,8.3,1.1,'C904 · 2 days / 15 instructional hours\nDr. Alfred Ang · v4.0 · 2 October 2026',21,GREY)
shape(s,9.1,1.7,3.6,3.6,TEAL,MSO_SHAPE.OVAL);text(s,9.25,2.9,3.3,1.1,'Lead with\nclarity',32,WHITE,True,PP_ALIGN.CENTER)
text(s,.65,6.3,12,.55,C.COURSE_URL,17,TEAL)
cards('Welcome and working agreements',['Introduce your role, a leadership challenge and one behaviour you want to practise.','Listen without interruption. Discuss actions and impacts; respect different perspectives.','Use the fictional scenarios for practice. Keep personal workplace details private.','Give specific peer feedback and record a next action after each exercise.'],'WELCOME')
cards('Your trainer and course materials',['Dr. Alfred Ang is the trainer named in the original C904 reference deck.','Open the current course page and the shared course folder. Each activity has its own complete pack.','Keep the Learner Guide, worksheet and checklist beside the slides. Detailed procedures are in the guide and activity files.'],'GET READY')
flow("How You'll Learn",['Understand','Observe','Practise','Reflect'],'Learn a framework, see an example, try a response with peers and choose an action to improve.')
cards('Learning outcomes',C.LEARNING_OUTCOMES,'WHAT YOU WILL LEARN')
cards('Two days, six connected topics',[f'Day 1: {C.DAY_THEMES[1]}. Topics 1–3; Activities 1–6.',f'Day 2: {C.DAY_THEMES[2]}. Topics 4–6; Activities 7–12.','Each day: 450 instructional minutes, 30 minutes of tea breaks and 60 minutes of lunch. Courseware timetable: 09:30–18:30.','Confirm the current run timing on the registration page. Use the Lesson Plan for the detailed teaching sequence.'],'LEARNING JOURNEY')
import importlib
acts=sum([getattr(importlib.import_module(f'data_domain{i}'),f'DOMAIN{i}') for i in range(1,7)],[])
for t in C.TOPICS:
 accent=accents[t['num']-1];divider(t['num'],t['title'],t['subtitle'])
 for j,o in enumerate(t['original']):
  ps=o['paragraphs']
  if any(term in o['title'].lower() for term in ['maslow','herzberg','self-determination']):theory_visual(o,accent)
  cards(o['title'],ps,f'TOPIC {t["num"]:02d} · CONCEPT LIBRARY',accent)
  if 'sandwich' in ' '.join(ps).lower(): cards('Specific feedback is clearer than a formula',['The reference deck introduces the sandwich approach. Praise used merely to soften criticism can make the message unclear.','Prefer a specific observed behaviour, its effect, the person’s perspective and an agreed next action. Give genuine appreciation separately when warranted.'],'TEACHING CLARIFICATION',accent)
  if 'Maslow' in o['title']:cards('Use motivation theories as discussion lenses',['People’s needs can overlap and vary with context. A hierarchy is not a diagnostic tool or a fixed sequence for every employee.','Ask about actual work conditions, resources and what matters to the person before choosing a support action.'],'TEACHING CLARIFICATION',accent)
 for lesson in t['lessons']:
  # diagrams carry relationships, with an original worked-example caption
  paragraphs=re.split(r'(?<=[.!?])\s+',lesson['text']);flow(lesson['title'],lesson['nodes'],' '.join(paragraphs[:2]),accent,kicker='APPLY THE FRAMEWORK')
  cards(lesson['title']+' in practice',[' '.join(paragraphs[2:])],f'WORKED EXAMPLE · {lesson["source"]}',accent)
  if lesson['title'].startswith('Weighted'):chart_matrix()
 for a in [a for a in acts if a['topic']==t['num']]:
  if a.get('original_practice_context'): cards('Original practice context',a['original_practice_context'].split('\n\n'),'OPTIONAL DISCUSSION',accent)
  s=base(f'Activity {a["num"]:02d}: {a["title"]}','PRACTISE · 35 MINUTES',accent)
  shape(s,.65,2.15,7.55,4.65,LIGHT);text(s,.95,2.45,6.95,1.7,a['desc'],18)
  text(s,.95,4.85,6.95,1.7,'Finished result: '+a['build'],18,accent,True)
  text(s,8.65,2.45,3.9,2.15,'Apply the concepts\nCompare perspectives\nUse the peer checklist',24,accent,True)
  text(s,8.65,5.1,3.9,1.6,'Detailed instructions and worksheets are in the activity folder and Learner Guide.',19,GREY)
 cards('Recap: '+t['title'],[C.LEARNING_OUTCOMES[t['num']-1],'Explain your choice with evidence, name a trade-off and record a review point.','Which behaviour will you keep practising after today?'],'TOPIC RECAP',accent)
cards('What You Achieved',C.LEARNING_OUTCOMES,'LEARNING OUTCOMES')
flow('Keep Practising',['Choose a behaviour','Try it at work','Ask for feedback','Review evidence'],'Use Activity 12 to plan reviews after 7, 14 and 30 days. Adapt the plan as you learn.')
for i in range(0,len(C.SOURCES),5):cards('Further reading',[name+'\n'+url for key,name,url in C.SOURCES[i:i+5]],'SOURCES · ACCESSED 2 OCTOBER 2026')
cards('Thank You',['Leadership Masterclass · C904 · v4.0','Continue using the worksheets and ask a colleague for specific feedback on your chosen behaviour.',C.COURSE_URL],'NEXT STEPS')
out=root/'courseware'/f'{C.SHORT_TITLE}-{C.VERSION}.pptx';prs.save(out);print('Saved',out,'slides',len(prs.slides))
