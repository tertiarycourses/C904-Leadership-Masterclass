from pathlib import Path
import sys,importlib
from docx import Document
from docx.shared import Pt,Inches
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
import html
root=Path(__file__).resolve().parent.parent;b=root/'.claude/skills/non-wsq-courseware-build/build';sys.path.insert(0,str(b));import course_data as C
acts=sum([getattr(importlib.import_module(f'data_domain{i}'),f'DOMAIN{i}') for i in range(1,7)],[])
styles=getSampleStyleSheet();styles['BodyText'].fontSize=10.5;styles['BodyText'].leading=15
styles['Heading1'].textColor=colors.HexColor('#087F8C')
def pdf(path,title,blocks):
 story=[Paragraph(html.escape(title),styles['Title']),Paragraph('Leadership Masterclass · C904 · v4.0 · Fictional practice material',styles['BodyText']),Spacer(1,16)]
 for kind,value in blocks:
  if kind=='h':story.append(Paragraph(html.escape(value),styles['Heading2']))
  else:story.append(Paragraph(html.escape(value).replace('\n','<br/>'),styles['BodyText']));story.append(Spacer(1,9))
 def foot(c,d):c.setFont('Helvetica',8);c.drawString(42,24,'C904 · Leadership Masterclass · Tertiary Infotech Academy');c.drawRightString(553,24,str(d.page))
 SimpleDocTemplate(str(path),pagesize=(595.28,841.89),rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42).build(story,onFirstPage=foot,onLaterPages=foot)
for a in acts:
 folder=root/a['folder']
 if a.get('original_practice_context'): pdf(folder/'original-practice-context.pdf','Original C904 practice context',[('p','Optional extension retained from the original reference deck. Apply the current framework, worksheet and checklist.'),('p',a['original_practice_context'])])
 for sample in [False,True]:
  stem='sample-completed-worksheet' if sample else 'worksheet';title=('Illustrative worked example: ' if sample else 'Worksheet: ')+a['title'];d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11);d.add_heading(title,0);d.add_paragraph('Leadership Masterclass · C904 · v4.0');d.add_paragraph('Name/group: ____________________  Date: ____________________');d.add_heading('Fictional scenario',1);d.add_paragraph(a['scenario'])
  blocks=[('h','Fictional scenario'),('p',a['scenario'])]
  for j,field in enumerate(a['fields']):
   value=a['sample'][j] if sample and j<len(a['sample']) else '\n________________________________________________\n________________________________________________\n________________________________________________'
   d.add_heading(field,2);d.add_paragraph(value);blocks.extend([('h',field),('p',value)])
  reflection='One improvement: name the owner, the evidence and the next review date.' if sample else 'What changed after peer feedback? ________________________________________\n\nMy next action and review date: __________________________________________'
  d.add_heading('Reflection and next action',1);d.add_paragraph(reflection);blocks.extend([('h','Reflection and next action'),('p',reflection)])
  if sample:d.add_paragraph('This is one possible response, not a model that must be copied. Other reasoned choices can be appropriate.');blocks.append(('p','This is one possible response. Other reasoned choices can be appropriate.'))
  d.save(folder/(stem+'.docx'));pdf(folder/(stem+'.pdf'),title,blocks)
 pdf(folder/'checklist.pdf','Practice checklist: '+a['title'],[('p','[  ] '+c) for c in a['checks']]+[('h','Peer observation'),('p','Observed words/actions: __________________________________________\n\nOne change for the next attempt: __________________________________')])
 guide=[('h','Goal'),('p',a['desc']),('h','What you will build'),('p',a['build']),('h','Fictional scenario'),('p',a['scenario']),('h','Preparation'),('p','Read the scenario. Open worksheet.docx or print worksheet.pdf. Work in pairs or triads. Use sample-completed-worksheet.pdf after making your own first attempt.')]
 for j,(s,cmd) in enumerate(a['steps'],1):guide.extend([('h',f'Step {j}'),('p',s)])
 guide.extend([('h','Test it'),('p',a['test']),('h','Debrief'),('p','What evidence changed your view? Which trade-off remains? Which behaviour will you practise next?')]);pdf(folder/'activity-guide.pdf',f'Activity {a["num"]:02d}: '+a['title'],guide)
print('Built',len(acts),'activity packs')
